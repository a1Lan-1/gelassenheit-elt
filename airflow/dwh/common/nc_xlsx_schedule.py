"""Parse Calltraffic xlsx: monthly timesheets + employee list."""

from __future__ import annotations

import io
import logging
import re
from datetime import date, datetime
from typing import Any

import openpyxl
import pandas as pd

logger = logging.getLogger(__name__)

_SKIP_DEFAULT = ("actual hours", "required hours (load plan)")
_WS_RE = re.compile(r"\s+", re.UNICODE)


def _norm_name(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).replace("\xa0", " ").strip()
    if not text:
        return None
    return _WS_RE.sub(" ", text)


def _to_date_str(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return None
    parsed = pd.to_datetime(text, dayfirst=True, errors="coerce")
    if pd.isna(parsed):
        return text[:10] if len(text) >= 10 else text
    return parsed.strftime("%Y-%m-%d")


def _to_float(value: object) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip().replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _cell(row: tuple[Any, ...], col_1based: int) -> Any:
    idx = col_1based - 1
    if idx < 0 or idx >= len(row):
        return None
    return row[idx]


def flatten_schedule_workbook(
    data: bytes,
    *,
    employees_sheet: str,
    plan_col_1based: int = 27,
    name_col_1based: int = 2,
    date_col_1based: int = 30,
    skip_name_labels: list[str] | None = None,
) -> pd.DataFrame:
    """Monthly sheets → date, employee, hours (plan > 0), same layout as GS AKC QUERY."""
    skip = {s.strip().lower() for s in (skip_name_labels or list(_SKIP_DEFAULT))}
    emp_sheet_l = employees_sheet.strip().lower()

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: list[dict[str, object]] = []
    used_sheets: list[str] = []
    try:
        for sheet_name in wb.sheetnames:
            if sheet_name.strip().lower() == emp_sheet_l:
                continue
            used_sheets.append(sheet_name)
            ws = wb[sheet_name]
            for row in ws.iter_rows(values_only=True):
                if not row:
                    continue
                name = _norm_name(_cell(row, name_col_1based))
                if not name:
                    continue
                if name.lower() in skip or name.lower().startswith("norm"):
                    continue
# employee lines: Col1 = number (int)
                seq = _cell(row, 1)
                if not isinstance(seq, (int, float)) or isinstance(seq, bool):
                    continue
                hours = _to_float(_cell(row, plan_col_1based))
                if hours is None or hours <= 0:
                    continue
                date_s = _to_date_str(_cell(row, date_col_1based))
                if not date_s:
                    continue
                out.append({"date": date_s, "employee": name, "hours": hours})
    finally:
        wb.close()

    df = pd.DataFrame(out, columns=["date", "employee", "hours"])
    logger.info("NC schedule flatten: %s rows, sheets=%s", len(df), used_sheets)
    return df


def parse_employees_sheet(
    data: bytes,
    *,
    sheet: str,
    header_map: dict[str, str],
    columns: list[str],
    required_column: str | None = "full_name",
) -> pd.DataFrame:
    """Employee list sheet → bronze columns."""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        # case-insensitive sheet match
        target = None
        for name in wb.sheetnames:
            if name.strip().lower() == sheet.strip().lower():
                target = name
                break
        if target is None:
            raise ValueError(f"Sheet {sheet!r} not found; available={wb.sheetnames}")
        ws = wb[target]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return pd.DataFrame(columns=columns)

        # map header cells → column index
        idx_by_bronze: dict[str, int] = {}
        for i, cell in enumerate(header_row or ()):
            if cell is None:
                continue
            key = str(cell).replace("\xa0", " ").strip()
            bronze = header_map.get(key)
            if bronze:
                idx_by_bronze[bronze] = i

        missing = [c for c in columns if c not in idx_by_bronze]
        if missing:
            logger.warning("NC employees missing header cols: %s (found=%s)", missing, list(idx_by_bronze))

        date_cols = {
            "training_start_date",
            "training_end_date",
            "line_start_date",
            "birthday",
            "dismissal_date",
        }
        out: list[dict[str, object]] = []
        for row in rows_iter:
            if not row or not any(c is not None and str(c).strip() for c in row):
                continue
            rec: dict[str, object] = {}
            for col in columns:
                idx = idx_by_bronze.get(col)
                raw = row[idx] if idx is not None and idx < len(row) else None
                if col == "full_name" or col == "supervisor":
                    rec[col] = _norm_name(raw)
                elif col in date_cols:
                    rec[col] = _to_date_str(raw)
                elif raw is None or (isinstance(raw, str) and not raw.strip()):
                    rec[col] = None
                else:
                    text = str(raw).replace("\xa0", " ").strip()
                    # non-binding placeholders
                    if text.lower() in ("optional", "opt", "-"):
                        rec[col] = None
                    else:
                        rec[col] = text
            if required_column and not rec.get(required_column):
                continue
            out.append(rec)
    finally:
        wb.close()

    df = pd.DataFrame(out, columns=columns)
    logger.info("NC employees parse: %s rows from %r", len(df), sheet)
    return df
